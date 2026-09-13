from pydantic import BaseModel,Field
from langchain.tools import tool
from docx import Document
import time
import pandas as pd
import io
import reportlab

#----------word工具----------
class WordParams(BaseModel):
    title:str = Field(...,description="文档标题")
    content:str=Field(...,description="文档内容")

@tool(args_schema=WordParams)
def word_tool(title:str,content:str)->str:
    """
    写入word文档
    """
    try:
        #1、新建一个文档对象
        doc=Document()
        #2、添加一个一级标题
        doc.add_heading(title,level=1)
        #3、添加一个段落
        doc.add_paragraph(content)
        #4、定义时间戳作为文件名字
        file_name=time.strftime("%Y%m%d%H%M%S",time.localtime())
        doc.save(f'{file_name}.docx')
        return "写入成功"
    except Exception as e:
        print(f"出现异常：{e}")
        return "出现异常"


#----------excel工具----------
class ExcelParams(BaseModel):
    title:str=Field(...,description="工作表名称")
    content:str=Field(...,description="表格数据")
@tool(args_schema=WordParams)
def excel_tool(title:str,content:str)->str:
    """
    写入excel文件，使用pandas将csv格式内容写入指定工作表
    第一行作为列名，后续行作为数据
    """
    try:
        #解析csv字符串作为DataFrame
        df = pd.read_csv(io.StringIO(content), skip_blank_lines=True)
        # 若内容为空，生成提示行
        if df.empty:
            df = pd.DataFrame([["（无内容）"]])
            # 生成时间戳文件名
        file_name = time.strftime("%Y%m%d%H%M%S", time.localtime())
        # 写入 Excel（使用 openpyxl 引擎，指定工作表名）
        with pd.ExcelWriter(f'{file_name}.xlsx', engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name=title[:31], index=False)
        return "Excel 文件写入成功"
    except Exception as e:
        print(f"Excel 异常：{e}")
        return "Excel 文件生成失败"


#----------pdf工具----------
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

class PdfParams(BaseModel):
    title: str = Field(..., description="PDF 标题（显示在页面顶部）")
    content: str = Field(..., description="PDF 正文内容，多段用换行分隔")

@tool(args_schema=PdfParams)
def pdf_tool(title: str, content: str) -> str:
    """写入 PDF 文件（.pdf），支持中文字体"""
    try:
        pdfmetrics.registerFont(TTFont('SimHei', 'C:/Windows/Fonts/simhei.ttf'))
        file_name = time.strftime("%Y%m%d%H%M%S", time.localtime())
        c = canvas.Canvas(f'{file_name}.pdf', pagesize=letter)
        c.setFont('SimHei', 16)
        width, height = letter

        # 标题使用中文字体
        c.setFont('SimHei', 16)
        c.drawString(1 * inch, height - 1 * inch, title)

        # 正文使用中文字体
        c.setFont('SimHei', 12)
        y = height - 1.5 * inch
        lines = content.split('\n')
        for line in lines:
            if y < 0.8 * inch:
                c.showPage()
                y = height - 1 * inch
                c.setFont('SimHei', 12)  # 新页面重新设置字体
            c.drawString(1 * inch, y, line[:100])  # 每行最多 100 字符
            y -= 0.3 * inch

        c.save()
        return f"PDF 文件写入成功：{file_name}.pdf"
    except Exception as e:
        print(f"PDF 异常：{e}")
        return "PDF 文件生成失败"

#----------txt工具----------
class TxtParams(BaseModel):
    title: str = Field(..., description="文本文件标题（写入首行）")
    content: str = Field(..., description="文本文件正文内容")

@tool(args_schema=TxtParams)
def txt_tool(title: str, content: str) -> str:
    """写入纯文本文件（.txt），标题作为第一行，内容接在后面"""
    try:
        file_name = time.strftime("%Y%m%d%H%M%S", time.localtime())
        with open(f'{file_name}.txt', 'w', encoding='utf-8') as f:
            f.write(title + '\n' + '=' * len(title) + '\n')  # 加下划线分隔
            f.write(content)
        return "TXT 文件写入成功"
    except Exception as e:
        print(f"TXT 异常：{e}")
        return "TXT 文件生成失败"


#测试
if __name__ == "__main__":
    print("=== 测试 Word ===")
    print(word_tool.invoke({"title": "工作报告", "content": "这是报告内容"}))

    print("\n=== 测试 Excel ===")
    print(excel_tool.invoke({
        "title": "人员清单",
        "content": "姓名,年龄,部门\n张三,28,研发\n李四,32,销售"
    }))

    print("\n=== 测试 PDF ===")
    print(pdf_tool.invoke({"title": "会议纪要", "content": "会议讨论了项目进度。\n下次会议定于周五。"}))

    print("\n=== 测试 TXT ===")
    print(txt_tool.invoke({"title": "备忘录", "content": "记得买牛奶。"}))



