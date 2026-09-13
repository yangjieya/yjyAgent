from pydantic import BaseModel,Field
from langchain.tools import tool
from docx import Document
from dotenv import load_dotenv
import os
import time

load_dotenv()

class WordParams(BaseModel):
    title:str= Field(...,description="文档标题")
    content:str= Field(...,description="文档内容")

@tool(args_schema=WordParams)
def word_tool(title:str,content:str)->str:
    """
     写入word文档
    """
    try:
        # 1. 创建一个新的文档对象
        doc = Document()
        # 2. 添加一级标题
        doc.add_heading(title, level=1)
        # 3. 添加一个普通段落
        p = doc.add_paragraph(content)
        #定义时间戳作为文件名
        file_name = time.strftime("%Y%m%d%H%M%S", time.localtime())
        #定义保存的目录位置
        file_path = f"{os.getenv('DOWNLOAD_DIR')}/{file_name}.docx"
        #保存路径
        doc.save(file_path)
        #定义下载链接
        url = f"http://localhost:8000/static/download/{file_name}.docx"
        return f"写入成功，word文档的链接地址:{url}"
    except Exception as e:
        print(f"出现异常{e}")
        return "出现异常"